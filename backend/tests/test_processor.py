"""
Unit and integration tests for processor.py.

Run from the backend/ directory:
    pytest tests/ -v
"""

import os
import sys

import numpy as np
import pandas as pd
import pytest

# Make sure the backend package is importable regardless of cwd
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from processor import (
    ACCOUNT_FA_GAIN_LOSS,
    ACCOUNT_INTERCO,
    INTERCO_CUSTOMER,
    build_group1,
    build_group2,
    build_je_item_text,
    build_output_df,
    find_excel_file,
    process_reclass,
)


# ── Fixtures ──────────────────────────────────────────────────────────────────

@pytest.fixture
def base_df():
    """Minimal DataFrame that mirrors post-merge, post-filter, post-JE-text Retirement_Data."""
    return pd.DataFrame({
        "Company Code":       [1000,                    1000,   2100],
        "Posting Date":       ["2025-01-15",            "2025-01-15",  "2025-01-15"],
        "JE Item Text":       ["reclass_gain_loss_4501","reclass_gain_loss_Manual sale", "reclass_gain_loss_4503"],
        "Profit Center":      [10000,  10000,  22345],
        "Costcenter":         [100000000, 100000004, 123456782],
        "Functional Area":    ["F620", "F630", "F630"],
        "Net Reclass Amount": [100,    200,    -500],
    })


@pytest.fixture
def excel_workbook(tmp_path):
    """Write a minimal two-sheet workbook; return (folder_path, file_path)."""
    retirement = pd.DataFrame({
        "Company Code":    [1000, 1000, 2100],
        "Posting Date":    ["2025-01-15", "2025-01-15", "2025-01-15"],
        "Cockpit Req No":  [4501.0, np.nan, 4503.0],
        "Item Text":       [np.nan, "Manual sale reclass 4005-Jan2025", np.nan],
        "Profit Center":   [10000, 10000, 22345],
        "Costcenter":      [100000000, 100000004, 123456782],
        "Functional Area": ["F620", "F630", "F630"],
        "Gain":            [100, 200, 0],
        "Loss":            [0, 0, -500],
    })
    cc_master = pd.DataFrame({
        "Costcenter":       [100000000, 100000004, 123456782],
        "Business Unit":    ["HDD Mfg",  "GSC",     "Corp Finance"],
        "Cost Center Name": ["CC01",     "CC03",    "CC05"],
        "Location":         ["Fremont",  "Woodlands","Penang"],
    })

    file_path = tmp_path / "reclass_test.xlsx"
    with pd.ExcelWriter(str(file_path), engine="openpyxl") as writer:
        retirement.to_excel(writer, sheet_name="Retirement_Data",   index=False)
        cc_master.to_excel(writer,  sheet_name="Cost_Center_Master", index=False)

    return str(tmp_path), str(file_path)


# ── find_excel_file ───────────────────────────────────────────────────────────

class TestFindExcelFile:
    def test_finds_xlsx(self, tmp_path):
        f = tmp_path / "data.xlsx"
        f.write_bytes(b"")
        assert find_excel_file(str(tmp_path)) == str(f)

    def test_raises_when_empty_folder(self, tmp_path):
        with pytest.raises(FileNotFoundError, match="No Excel file found"):
            find_excel_file(str(tmp_path))

    def test_raises_when_no_excel_files(self, tmp_path):
        (tmp_path / "notes.txt").write_text("hello")
        with pytest.raises(FileNotFoundError):
            find_excel_file(str(tmp_path))


# ── build_je_item_text ────────────────────────────────────────────────────────

class TestBuildJeItemText:
    def _make_df(self, req_no, item_text):
        return pd.DataFrame({
            "Cockpit Req No": [req_no],
            "Item Text":      [item_text],
        })

    def test_uses_cockpit_req_no_when_present(self):
        df = self._make_df(4501.0, np.nan)
        result = build_je_item_text(df)
        assert result.iloc[0] == "reclass_gain_loss_4501"

    def test_falls_back_to_item_text_when_req_no_missing(self):
        df = self._make_df(np.nan, "Manual sale reclass 4005-Jan2025")
        result = build_je_item_text(df)
        assert result.iloc[0] == "reclass_gain_loss_Manual sale reclass 4005-Jan2025"

    def test_truncated_to_50_chars(self):
        long_text = "A" * 60
        df = self._make_df(np.nan, long_text)
        result = build_je_item_text(df)
        assert len(result.iloc[0]) == 50

    def test_empty_item_text_fallback(self):
        df = self._make_df(np.nan, np.nan)
        result = build_je_item_text(df)
        assert result.iloc[0] == "reclass_gain_loss_"

    def test_preserves_dataframe_index(self):
        # After filtering, df index may not start at 0
        df = pd.DataFrame({
            "Cockpit Req No": [4501.0, np.nan],
            "Item Text":      [np.nan,  "text"],
        }, index=[5, 10])
        result = build_je_item_text(df)
        assert list(result.index) == [5, 10]


# ── build_group1 ──────────────────────────────────────────────────────────────

class TestBuildGroup1:
    def test_account_number_is_7615000(self, base_df):
        g = build_group1(base_df)
        assert (g["account_number"] == ACCOUNT_FA_GAIN_LOSS).all()

    def test_positive_amount_goes_to_debit(self, base_df):
        g = build_group1(base_df)
        positive_rows = g[g["Net Reclass Amount"] > 0]
        assert (positive_rows["Transaction Currency Debit"] > 0).all()
        assert (positive_rows["Transaction Currency Credit"] == 0).all()

    def test_negative_amount_goes_to_credit(self, base_df):
        g = build_group1(base_df)
        negative_rows = g[g["Net Reclass Amount"] < 0]
        assert (negative_rows["Transaction Currency Credit"] < 0).all()
        assert (negative_rows["Transaction Currency Debit"] == 0).all()

    def test_costcenter_is_preserved(self, base_df):
        g = build_group1(base_df)
        assert "Costcenter" in g.columns
        # Cost centers from base_df should appear in group1
        assert set(g["Costcenter"]).issubset(set(base_df["Costcenter"]))

    def test_sums_within_group_keys(self):
        df = pd.DataFrame({
            "Company Code":       [1000, 1000],
            "Posting Date":       ["2025-01-15", "2025-01-15"],
            "JE Item Text":       ["reclass_gain_loss_4501", "reclass_gain_loss_4501"],
            "Profit Center":      [10000, 10000],
            "Costcenter":         [100000000, 100000000],
            "Functional Area":    ["F620", "F620"],
            "Net Reclass Amount": [100, 150],
        })
        g = build_group1(df)
        assert len(g) == 1
        assert g.iloc[0]["Net Reclass Amount"] == 250


# ── build_group2 ──────────────────────────────────────────────────────────────

class TestBuildGroup2:
    def test_account_number_is_6645000(self, base_df):
        g = build_group2(base_df)
        assert (g["account_number"] == ACCOUNT_INTERCO).all()

    def test_costcenter_is_blank(self, base_df):
        g = build_group2(base_df)
        assert (g["Costcenter"] == "").all()

    def test_positive_amount_goes_to_credit(self, base_df):
        """Group2 has reversed sign logic vs Group1."""
        g = build_group2(base_df)
        positive_rows = g[g["Net Reclass Amount"] > 0]
        assert (positive_rows["Transaction Currency Credit"] > 0).all()
        assert (positive_rows["Transaction Currency Debit"] == 0).all()

    def test_negative_amount_goes_to_debit(self, base_df):
        g = build_group2(base_df)
        negative_rows = g[g["Net Reclass Amount"] < 0]
        assert (negative_rows["Transaction Currency Debit"] < 0).all()
        assert (negative_rows["Transaction Currency Credit"] == 0).all()

    def test_groups_without_costcenter(self, base_df):
        """Two rows with different cost centers but same PC should merge into one."""
        df = pd.DataFrame({
            "Company Code":       [1000, 1000],
            "Posting Date":       ["2025-01-15", "2025-01-15"],
            "JE Item Text":       ["reclass_gain_loss_4501", "reclass_gain_loss_4501"],
            "Profit Center":      [10000, 10000],
            "Costcenter":         [100000000, 100000004],  # different CC, same PC
            "Functional Area":    ["F620", "F620"],
            "Net Reclass Amount": [100, 200],
        })
        g = build_group2(df)
        assert len(g) == 1
        assert g.iloc[0]["Net Reclass Amount"] == 300


# ── build_output_df ───────────────────────────────────────────────────────────

class TestBuildOutputDf:
    @pytest.fixture
    def combined_df(self, base_df):
        import pandas as pd
        g1 = build_group1(base_df)
        g2 = build_group2(base_df)
        return pd.concat([g1, g2], ignore_index=True)

    def test_all_required_columns_present(self, combined_df):
        out = build_output_df(combined_df)
        expected_cols = [
            "*Company Code (4)", "*Journal Entry Type (2)", "*Journal Entry Date",
            "*Posting Date", "*Transaction Currency (3)",
            "*IC Service Transaction Type (5)", "Document Header Text (25)",
            "Ledger Group (4)", "Company Code (4)", "G/L Account (10)",
            "Transaction Currency Debit", "Transaction Currency Credit",
            "Item Text (50)", "Amount in Company Code Currency",
            "Cost Center (10)", "Profit Center (10)", "Trading Partner (6)",
            "Customer (10)", "Product number (40)",
        ]
        for col in expected_cols:
            assert col in out.columns, f"Missing column: {col}"

    def test_customer_set_only_for_interco_account(self, combined_df):
        out = build_output_df(combined_df)
        interco_rows = out[out["G/L Account (10)"] == ACCOUNT_INTERCO]
        other_rows   = out[out["G/L Account (10)"] != ACCOUNT_INTERCO]
        assert (interco_rows["Customer (10)"] == INTERCO_CUSTOMER).all()
        assert (other_rows["Customer (10)"] == "").all()

    def test_product_number_set_only_for_interco_account(self, combined_df):
        out = build_output_df(combined_df)
        interco_rows = out[out["G/L Account (10)"] == ACCOUNT_INTERCO]
        other_rows   = out[out["G/L Account (10)"] != ACCOUNT_INTERCO]
        assert interco_rows["Product number (40)"].str.startswith("PC_").all()
        assert (other_rows["Product number (40)"] == "").all()

    def test_debit_credit_are_absolute_values(self, combined_df):
        out = build_output_df(combined_df)
        assert (out["Transaction Currency Debit"]  >= 0).all()
        assert (out["Transaction Currency Credit"] >= 0).all()

    def test_journal_entry_type_is_SA(self, combined_df):
        out = build_output_df(combined_df)
        assert (out["*Journal Entry Type (2)"] == "SA").all()


# ── process_reclass (integration) ────────────────────────────────────────────

class TestProcessReclass:
    def test_creates_output_file(self, excel_workbook, tmp_path):
        input_folder, _ = excel_workbook
        output_folder = str(tmp_path / "output")

        output_path, output_filename = process_reclass(input_folder, output_folder)

        assert os.path.exists(output_path)
        assert output_filename.startswith("JE_")
        assert output_filename.endswith(".xlsx")

    def test_output_has_correct_sheet(self, excel_workbook, tmp_path):
        input_folder, _ = excel_workbook
        output_folder = str(tmp_path / "output")

        output_path, _ = process_reclass(input_folder, output_folder)
        result = pd.read_excel(output_path, sheet_name="JE_Output")

        assert not result.empty
        assert "*Company Code (4)" in result.columns
        assert "G/L Account (10)" in result.columns

    def test_output_contains_both_accounts(self, excel_workbook, tmp_path):
        input_folder, _ = excel_workbook
        output_folder = str(tmp_path / "output")

        output_path, _ = process_reclass(input_folder, output_folder)
        result = pd.read_excel(output_path, sheet_name="JE_Output")

        accounts = set(result["G/L Account (10)"].unique())
        assert ACCOUNT_FA_GAIN_LOSS in accounts
        assert ACCOUNT_INTERCO in accounts

    def test_header_has_green_fill(self, excel_workbook, tmp_path):
        from openpyxl import load_workbook

        input_folder, _ = excel_workbook
        output_folder = str(tmp_path / "output")

        output_path, _ = process_reclass(input_folder, output_folder)
        wb = load_workbook(output_path)
        ws = wb["JE_Output"]

        first_cell = ws.cell(row=1, column=1)
        assert first_cell.fill.fgColor.rgb.upper().endswith("90EE90")

    def test_raises_on_missing_input_folder(self, tmp_path):
        with pytest.raises(FileNotFoundError):
            process_reclass(str(tmp_path / "nonexistent"), str(tmp_path / "out"))

    def test_raises_on_no_excel_in_folder(self, tmp_path):
        (tmp_path / "notes.txt").write_text("hello")
        with pytest.raises(FileNotFoundError):
            process_reclass(str(tmp_path), str(tmp_path / "out"))

    def test_creates_output_folder_if_missing(self, excel_workbook, tmp_path):
        input_folder, _ = excel_workbook
        output_folder = str(tmp_path / "deep" / "nested" / "output")

        process_reclass(input_folder, output_folder)

        assert os.path.isdir(output_folder)
